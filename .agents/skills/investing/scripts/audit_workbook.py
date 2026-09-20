"""Check a synchronized investing snapshot; visual review stays a separate gate.

Run from the checkout: uv run python <this-script> <snapshot/model.json> --recalculate
"""
import argparse
import json
from pathlib import Path

import openpyxl

from investing_research.excel import read_inputs, validate_snapshot_live, verify_snapshot


def audit(path, recalculate=False):
    result = verify_snapshot(path)
    workbook = path.parent / "model.xlsx"
    data, layout, _, _ = read_inputs(workbook)
    issues = []
    links = comments = 0
    book = openpyxl.load_workbook(workbook)
    cached = openpyxl.load_workbook(workbook, data_only=True)
    try:
        visible = [sheet.title for sheet in book if sheet.sheet_state == "visible"]
        for row, fact in enumerate(data["historical"], 7):
            cell = book["Historical"].cell(row, 8)
            if fact["source"].startswith(("http://", "https://")):
                link = cell.hyperlink
                actual = (link.target or "") if link else ""
                if link and link.location:
                    actual += "#" + link.location
                if actual != fact["source"]:
                    issues.append(f"Historical!H{row}: source hyperlink differs from recorded URL")
                links += 1
        # Comments retain original source provenance even when a user overrides an input.
        provenance = layout["baseline"].get("provenance", {})
        for pointer, location in layout["inputs"].items():
            if pointer not in provenance:
                continue
            comment = book[location["sheet"]][location["cell"]].comment
            expected = "\n".join(f"{k}: {v}" for k, v in provenance[pointer].items())
            if comment is None or comment.text != expected:
                issues.append(f"{location['sheet']}!{location['cell']}: source comment differs")
            comments += 1
        if book["_Model"].sheet_state != "hidden":
            issues.append("Model manifest must stay hidden")
        for sheet in cached:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "e":
                        issues.append(f"{sheet.title}!{cell.coordinate}: {cell.value}")
        for row in range(5, 11):
            if cached["Checks"].cell(row, 2).value != "PASS":
                issues.append(f"Checks!B{row}: expected PASS")
        if "Analysis" in cached:
            for cell in ("B35", "B46"):
                value = cached["Analysis"][cell].value
                if not isinstance(value, (int, float)) or abs(value) > 1e-8:
                    issues.append(f"Analysis!{cell}: bridge must reconcile to zero")
            if len(book["Analysis"]._charts) != 2:
                issues.append("Analysis: expected both historical bridge charts")
        if len(book["Summary"]._charts) != 2:
            issues.append("Summary: expected both model charts")
    finally:
        book.close()
        cached.close()
    live = validate_snapshot_live(path) if recalculate else {"status": "not checked this run"}
    return {
        "workbook_sha256": result["workbook_sha256"],
        "structural_status": "failed" if issues else "passed",
        "presentation_version": layout.get("presentation_version", 1),
        "visible_sheets": visible,
        "formulas": len(layout["formula_cells"]),
        "historical_rows": len(data["historical"]),
        "source_links_checked": links,
        "source_comments_checked": comments,
        "live_recalculation": live,
        "visual_status": "pending separate sheet-by-sheet and print inspection",
        "source_accessibility": "not checked; preserved URLs compared only",
        "issues": issues,
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("snapshot", type=Path)
    parser.add_argument("--recalculate", action="store_true")
    args = parser.parse_args()
    report = audit(args.snapshot.resolve(), args.recalculate)
    print(json.dumps(report, indent=2))
    raise SystemExit(bool(report["issues"]))
