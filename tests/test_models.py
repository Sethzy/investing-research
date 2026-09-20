import copy
import json
from pathlib import Path
import pytest
from investing_research.models import analyse, run_model, scenario_from_base

ROOT = Path(__file__).resolve().parents[1]


def fixture(name="mine"):
    return json.loads((ROOT / f"examples/model-{name}.json").read_text())


def test_finite_mine_independent_values():
    result = analyse(fixture())
    base = result["scenarios"]["base"]
    # 1,000 sales minus 400 costs minus 150 tax minus 100 capex = 350;
    # 50% ownership: 175, then closure reduces second year to 150.
    assert [x["fcf"] for x in base["years"]] == [175, 150]
    assert base["enterprise_value"] == pytest.approx(283.0578512396694)
    assert base["equity_value"] == pytest.approx(363.0578512396694)
    assert base["per_share"] == pytest.approx(3.630578512396694)
    assert base["minimum_illustrative_cash"] == 100
    assert result["reverse_valuation"]["status"] == "solved"


def test_offline_fcf_known_value():
    base = analyse(fixture("fcf"))["scenarios"]["base"]
    assert [x["fcf"] for x in base["years"]] == [80, 80]
    assert base["enterprise_value"] == pytest.approx(138.84297520661156)


def test_scenario_isolated_and_all_years():
    original = fixture()
    before = copy.deepcopy(original)
    revised = scenario_from_base(original, 0.8, 1.1)
    assert original == before
    assert revised["scenarios"]["base"]["years"][0]["price"] == 80
    assert revised["scenarios"]["base"]["years"][1]["unit_cost"] == 44
    # 800 revenue - 440 cost - 90 tax - 100 capex, half attributable.
    assert analyse(revised)["scenarios"]["base"]["years"][0]["fcf"] == 85


@pytest.mark.parametrize(
    "mutation",
    [
        lambda d: d.update(diluted_shares=0),
        lambda d: d["scenarios"]["base"].update(ownership_basis="attributable"),
        lambda d: d["scenarios"]["base"]["years"][0].update(fx=0),
        lambda d: d["scenarios"]["base"]["years"][1].update(year=2030),
        lambda d: d["provenance"].pop("/cash"),
        lambda d: d.update(debt=float("nan")),
        lambda d: d.update(sector="bank"),
    ],
)
def test_invalid_inputs_block(mutation):
    data = fixture()
    mutation(data)
    with pytest.raises(ValueError):
        analyse(data)


def test_workbook_recalculates_and_immutable_run(tmp_path):
    result = run_model(ROOT / "examples/model-mine.json", tmp_path)
    directory = Path(result["output_dir"])
    assert (directory / "report.md").exists()
    assert (directory / "cash-flows.svg").exists()
    import shutil

    if shutil.which("soffice"):
        assert result["workbook_validation"]["status"] == "passed"
    else:
        assert result["workbook_validation"]["status"] == "unverified"
    snapshot = (directory / "inputs.json").read_bytes()
    repeated = run_model(ROOT / "examples/model-mine.json", tmp_path)
    assert repeated == result
    assert snapshot == (directory / "inputs.json").read_bytes()
    import openpyxl

    book = openpyxl.load_workbook(directory / "model.xlsx", data_only=False)
    assert book["Summary"]["B2"].data_type == "f"
    assert book["base"]["B2"].value == 100
    book.close()


def test_editing_workbook_input_recalculates_formulas(tmp_path):
    import shutil

    if not shutil.which("soffice"):
        pytest.skip("Independent spreadsheet engine unavailable")
    import openpyxl
    from investing_research.workbook import write_workbook, verify_workbook

    data = fixture()
    path = tmp_path / "editable.xlsx"
    write_workbook(data, analyse(data), path)
    book = openpyxl.load_workbook(path)
    book["base"]["C12"] = 80  # first year commodity price only
    book.save(path)
    book.close()
    data["scenarios"]["base"]["years"][0]["price"] = 80
    expected = analyse(data)
    assert expected["scenarios"]["base"]["years"][0]["fcf"] == 100
    assert verify_workbook(path, expected)["status"] == "passed"


@pytest.mark.parametrize("omit", [False, True])
def test_missing_quote_allows_valuation_but_no_reverse(omit):
    data = fixture()
    for key in ("reference_price", "quote_date"):
        if omit:
            data.pop(key)
        else:
            data[key] = None
    result = analyse(data)
    assert result["reference_price"] is None
    assert result["quote_date"] is None
    assert result["scenarios"]["base"]["per_share"] == pytest.approx(3.630578512396694)
    assert result["reverse_valuation"]["status"] == "unavailable"
    assert "No verified reference price" in result["reverse_valuation"]["reason"]
    assert len(result["driver_comparisons"]) == 12


def test_quote_pair_zero_unknown_and_synthetic_integrity():
    for modification in (
        {"quote_date": None},
        {"reference_price": 0},
        {"synthetic": False},
        {"terminal_growth": 0.02},
    ):
        data = fixture()
        data.update(modification)
        with pytest.raises(ValueError):
            analyse(data)


def test_missing_quote_workbook_and_report(tmp_path):
    data = fixture()
    data["reference_price"] = data["quote_date"] = None
    source = tmp_path / "input.json"
    source.write_text(json.dumps(data))
    result = run_model(source, tmp_path / "out")
    directory = Path(result["output_dir"])
    import openpyxl

    workbook = openpyxl.load_workbook(directory / "model.xlsx")
    assert workbook["Summary"]["B7"].value is None
    workbook.close()
    assert "unavailable — valuation only" in (directory / "report.md").read_text()


def test_production_proxy_requires_illustrative_label():
    data = fixture()
    data["scenarios"]["base"]["volume_basis"] = "imputed_production"
    with pytest.raises(ValueError, match="illustrative"):
        analyse(data)
    data["purpose"] = "illustrative_sensitivity"
    assert analyse(data)["purpose"] == "illustrative_sensitivity"


def test_physical_history_has_no_currency_but_monetary_does():
    data = fixture()
    fact = dict(
        entity="example",
        metric="production",
        value=10,
        unit="tonnes",
        currency=None,
        period="2025",
        publication_date="2026-01-01",
        source="synthetic://fixture",
        page="1",
        basis="reported",
        ownership_basis="whole_operation",
    )
    data["historical"] = [fact]
    analyse(data)
    fact["unit"] = "AUD"
    with pytest.raises(ValueError, match="currency"):
        analyse(data)


def test_workbook_revalidation_after_canonical_input_reload(tmp_path):
    import shutil

    if not shutil.which("soffice"):
        pytest.skip("Independent spreadsheet engine unavailable")
    from investing_research.workbook import verify_workbook

    result = run_model(ROOT / "examples/model-mine.json", tmp_path)
    directory = Path(result["output_dir"])
    reloaded = json.loads((directory / "inputs.json").read_text())
    expected = analyse(reloaded)
    assert list(result["scenarios"]) != list(expected["scenarios"])
    assert verify_workbook(directory / "model.xlsx", expected)["status"] == "passed"


@pytest.mark.parametrize("label", ["base", None, "unknown"])
def test_workbook_rejects_duplicate_or_missing_scenario_labels(tmp_path, label):
    import shutil

    if not shutil.which("soffice"):
        pytest.skip("Independent spreadsheet engine unavailable")
    import openpyxl
    from investing_research.workbook import write_workbook, verify_workbook

    data = fixture()
    expected = analyse(data)
    path = tmp_path / "labels.xlsx"
    write_workbook(data, expected, path)
    workbook = openpyxl.load_workbook(path)
    workbook["Summary"]["A2"] = label
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="scenario label"):
        verify_workbook(path, expected)


@pytest.mark.parametrize("failure", [RuntimeError, KeyboardInterrupt])
def test_model_failure_preserves_stage_and_retry_publishes_atomically(tmp_path, monkeypatch, failure):
    from investing_research import workbook

    def interrupt(*args):
        raise failure("simulated interrupted workbook")

    source = ROOT / "examples/model-mine.json"
    with monkeypatch.context() as patch:
        patch.setattr(workbook, "write_workbook", interrupt)
        with pytest.raises(failure):
            run_model(source, tmp_path)
    assert [path.name for path in tmp_path.iterdir()] == [".incomplete"]
    staged_inputs = list((tmp_path / ".incomplete").glob("*/inputs.json"))
    assert len(staged_inputs) == 1
    snapshot = staged_inputs[0].read_bytes()
    result = run_model(source, tmp_path)
    final = Path(result["output_dir"])
    assert final.parent == tmp_path
    assert final.name == result["run_id"]
    assert (final / "model.json").is_file()
    assert json.loads((final / "model.json").read_text())["output_dir"] == str(final)
    assert staged_inputs[0].read_bytes() == snapshot


def test_preexisting_incomplete_hash_run_preserved_and_retried(tmp_path):
    import hashlib
    from investing_research.models import VERSION

    data = fixture()
    canonical = json.dumps(data, sort_keys=True, ensure_ascii=False, allow_nan=False).encode()
    digest = hashlib.sha256(canonical + VERSION.encode()).hexdigest()[:16]
    legacy = tmp_path / digest
    legacy.mkdir()
    (legacy / "inputs.json").write_bytes(canonical)
    (legacy / "interrupted.txt").write_text("keep this prior attempt")
    result = run_model(ROOT / "examples/model-mine.json", tmp_path)
    assert result["run_id"] == digest
    preserved = list((tmp_path / ".incomplete").glob("*/interrupted.txt"))
    assert len(preserved) == 1
    assert preserved[0].read_text() == "keep this prior attempt"
    assert (legacy / "model.json").is_file()
    assert not (legacy / "interrupted.txt").exists()
