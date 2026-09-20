import copy
import json
import shutil
from pathlib import Path

import openpyxl
import pytest
from typer.testing import CliRunner

from investing_research.cli import app
from investing_research.contracts import Company, Settings
from investing_research.excel import create, file_hash, read_inputs, refresh, status, synchronize, verify_snapshot
from investing_research.workspace import Workspace, read_json, write_json

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def model(tmp_path):
    workbook = tmp_path / "working.xlsx"
    create(ROOT / "examples/model-mine.json", workbook)
    return workbook


def edit(path, pointer, value):
    _, manifest, _, _ = read_inputs(path)
    book = openpyxl.load_workbook(path)
    location = manifest["inputs"][pointer]
    book[location["sheet"]][location["cell"]] = value
    book.save(path)
    book.close()


def test_edit_inputs_keeps_source_and_marks_provenance(model):
    source = json.loads((ROOT / "examples/model-mine.json").read_text())
    edit(model, "/scenarios/base/years/0/price", 80)
    inputs, _, _, changes = read_inputs(model)
    assert inputs["scenarios"]["base"]["years"][0]["price"] == 80
    assert changes[0]["before"] == 100
    assert inputs["provenance"][changes[0]["path"]]["author"] == "workbook user"
    assert json.loads((ROOT / "examples/model-mine.json").read_text()) == source


@pytest.mark.parametrize("cell,value", [("B11", 999), ("B11", '=WEBSERVICE("https://example.com")'), ("A5", "Wrong company")])
def test_formula_or_structural_changes_rejected(model, cell, value):
    book = openpyxl.load_workbook(model)
    book["Valuation"][cell] = value
    book.save(model)
    book.close()
    with pytest.raises(ValueError, match="changed"):
        read_inputs(model)


def test_bad_inputs_and_case_rejected(model):
    edit(model, "/diluted_shares", 0)
    with pytest.raises(ValueError, match="positive"):
        read_inputs(model)


def test_refresh_preserves_overrides_across_repeated_refreshes(model, tmp_path):
    edit(model, "/scenarios/base/years/0/price", 80)
    incoming = json.loads((ROOT / "examples/model-mine.json").read_text())
    incoming["cash"] = 125
    incoming["scenarios"]["base"]["years"][0]["price"] = 95
    source = tmp_path / "new.json"
    write_json(source, incoming)
    before = file_hash(model)
    output = tmp_path / "refreshed.xlsx"
    result = refresh(model, source, output)
    inputs = read_inputs(output)[0]
    assert inputs["cash"] == 125
    assert inputs["scenarios"]["base"]["years"][0]["price"] == 80
    assert result["conflicts"][0]["new_source"] == 95
    assert file_hash(model) == before
    incoming["scenarios"]["base"]["years"][0]["price"] = 90
    write_json(source, incoming)
    second = tmp_path / "second.xlsx"
    refresh(output, source, second)
    assert read_inputs(second)[0]["scenarios"]["base"]["years"][0]["price"] == 80


def test_create_does_not_overwrite(model):
    before = model.read_bytes()
    with pytest.raises(ValueError, match="exists"):
        create(ROOT / "examples/model-fcf.json", model)
    assert model.read_bytes() == before


def test_refresh_keeps_quote_pair_and_nulls_across_repeated_updates(model, tmp_path):
    edit(model, "/reference_price", 3)
    edit(model, "/quote_date", "2026-09-19")
    source = ROOT / "examples/model-mine.json"
    first, second = tmp_path / "one.xlsx", tmp_path / "two.xlsx"
    refresh(model, source, first)
    refresh(first, source, second)
    data = read_inputs(second)[0]
    assert (data["reference_price"], data["quote_date"]) == (3, "2026-09-19")
    edit(second, "/reference_price", None)
    # Complete the paired edit before validation (intermediate state is intentionally invalid).
    book = openpyxl.load_workbook(second)
    from investing_research.excel import manifest
    location = manifest(book)["inputs"]["/quote_date"]
    book[location["sheet"]][location["cell"]] = None
    book.save(second)
    book.close()
    third, fourth = tmp_path / "three.xlsx", tmp_path / "four.xlsx"
    refresh(second, source, third)
    refresh(third, source, fourth)
    assert read_inputs(fourth)[0]["reference_price"] is None
    assert read_inputs(fourth)[0]["quote_date"] is None


def test_missing_engine_never_publishes_cached_values(model, tmp_path, monkeypatch):
    monkeypatch.setattr("investing_research.excel.shutil.which", lambda _: None)
    with pytest.raises(ValueError, match="cached"):
        synchronize(Workspace(tmp_path), model)
    assert not list(tmp_path.glob("companies/*/models/*/model.json"))
    assert not list(tmp_path.glob("state/excel/*.json"))


@pytest.mark.skipif(not shutil.which("soffice"), reason="Independent spreadsheet engine unavailable")
def test_edit_sync_report_dossier_and_staleness(model, tmp_path):
    ws = Workspace(tmp_path)
    company = Company(id="demo-mine", name="Synthetic Mine", exchange="DEMO", ticker="MINE", currency="AUD", reporting_currency="AUD", sector="mining", website="https://example.com")
    write_json(ws.settings_path, Settings(companies=[company]).model_dump(mode="json"))
    edit(model, "/scenarios/base/years/0/price", 80)
    result = synchronize(ws, model)
    # Independent hand calculation: 800 revenue - 400 cost - 100 tax - 100 capex,
    # multiplied by 50% ownership gives 100 in year one; year two remains 150.
    expected = (100 / 1.1 + 150 / 1.1**2 + 100 - 20) / 100
    assert result["scenarios"]["base"]["per_share"] == pytest.approx(expected)
    assert result["authority"] == "excel"
    folder = Path(result["output_dir"])
    assert f"{expected:.4f}" in (folder / "report.md").read_text()
    assert "[Open the detailed Excel model](model.xlsx)" in (folder / "report.md").read_text()
    assert "Detailed Excel snapshot" in (tmp_path / "companies/demo-mine/dossier.md").read_text()
    assert status(ws)[0]["status"] == "current"
    assert verify_snapshot(folder / "model.json")["authority"] == "excel"
    frozen = file_hash(folder / "model.xlsx")
    edit(model, "/scenarios/base/years/1/price", 80)
    assert status(ws)[0]["status"] == "unsynchronized_edits"
    assert file_hash(folder / "model.xlsx") == frozen
    from investing_research.reports import dossier
    assert "not synchronized" in dossier(ws, company.id).read_text()
    second = synchronize(ws, model)
    assert second["run_id"] != result["run_id"]
    assert file_hash(folder / "model.xlsx") == frozen


@pytest.mark.skipif(not shutil.which("soffice"), reason="Independent spreadsheet engine unavailable")
def test_bad_snapshot_result_and_changed_workbook_rejected(model, tmp_path):
    result = synchronize(Workspace(tmp_path), model)
    folder = Path(result["output_dir"])
    original = copy.deepcopy(result)
    result["scenarios"]["base"]["per_share"] = 99
    write_json(folder / "model.json", result)
    with pytest.raises(ValueError, match="reproduce"):
        verify_snapshot(folder / "model.json")
    write_json(folder / "model.json", original)
    edit(folder / "model.xlsx", "/cash", 999)
    with pytest.raises(ValueError, match="edited"):
        verify_snapshot(folder / "model.json")
    with pytest.raises(ValueError, match="edited"):
        synchronize(Workspace(tmp_path), model)


@pytest.mark.skipif(not shutil.which("soffice"), reason="Independent spreadsheet engine unavailable")
def test_reuse_missing_snapshot_report_rejected(model, tmp_path):
    result = synchronize(Workspace(tmp_path), model)
    (Path(result["output_dir"]) / "report.md").unlink()
    with pytest.raises(ValueError, match="incomplete"):
        synchronize(Workspace(tmp_path), model)


def test_cli_create_and_invalid_selector(model, tmp_path):
    result = CliRunner().invoke(app, ["--root", str(tmp_path), "excel-create", str(ROOT / "examples/model-mine.json"), "--output", "new.xlsx"])
    # Root confines the workspace; caller should copy supplied inputs into it.
    assert result.exit_code != 0
    book = openpyxl.load_workbook(model)
    book["Assumptions"]["B4"] = "invented"
    book.save(model)
    book.close()
    with pytest.raises(ValueError, match="Select"):
        read_inputs(model)


def test_mlx_history_reconciles():
    data = read_json(ROOT / "examples/model-mlx.json")
    assert data["purpose"] == "illustrative_sensitivity"
    assert len(data["historical"]) >= 73
    assert all("/Users/" not in str(row) for row in data["historical"])

    history = {(r["period"], r["metric"]): r["value"] for r in data["historical"]}
    for year in (2024, 2025):
        fy, date = f"FY{year}", f"{year}-12-31"
        assert history[fy, "revenue"] - history[fy, "cost_of_sales"] == history[fy, "gross_profit"]
        assert history[fy, "profit_before_tax"] - history[fy, "income_tax_expense"] == history[fy, "net_income"]
        assert history[date, "total_assets"] == history[date, "total_liabilities"] + history[date, "total_equity"]
        assert history[date, "total_assets"] == history[date, "total_current_assets"] + history[date, "total_noncurrent_assets"]
        assert history[fy, "opening_cash_and_cash_equivalents"] + history[fy, "net_change_in_cash"] == history[date, "cash_and_cash_equivalents"]


def test_model_cli_without_engine_creates_editable_workbook(tmp_path, monkeypatch):
    shutil.copy2(ROOT / "examples/model-mine.json", tmp_path / "inputs.json")
    monkeypatch.setattr("investing_research.cli.shutil.which", lambda _: None)
    result = CliRunner().invoke(app, ["--root", str(tmp_path), "model", "inputs.json"])
    assert result.exit_code == 0, result.output
    payload = json.loads(result.output)
    assert payload["status"] == "editable_unverified"
    working = Path(payload["workbook"])
    edit(working, "/scenarios/base/years/0/price", 80)
    rerun = CliRunner().invoke(app, ["--root", str(tmp_path), "model", "inputs.json"])
    assert rerun.exit_code == 0, rerun.output
    assert read_inputs(working)[0]["scenarios"]["base"]["years"][0]["price"] == 80
    assert not list(tmp_path.glob("companies/*/models/*/report.md"))
