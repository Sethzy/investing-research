import copy
import json
import shutil
from pathlib import Path

import openpyxl
import pytest

from investing_research.excel import read_inputs, recalculate, refresh
from investing_research.excel_layout import write_excel_model

ROOT = Path(__file__).resolve().parents[1]


def inputs():
    data = json.loads((ROOT / 'examples/mlx/editions/2026-09-20-analysis/inputs.json').read_text())
    template = next(f for f in data['historical'] if f['metric'] == 'net_operating_cash_flow' and f['period'] == 'H1 2026')
    for metric, value in [('financial_assets_fvtpl', 48417000), ('investment_in_associates', 43864000)]:
        data['historical'].append(dict(template, metric=metric, value=value))
    return data


def test_decision_arithmetic_no_double_count_and_roundtrip(tmp_path):
    if not shutil.which('soffice'):
        pytest.skip('LibreOffice needed')
    data = inputs()
    path = tmp_path / 'decision.xlsx'
    write_excel_model(data, path, presentation_version=4)
    actual, layout, _, _ = read_inputs(path)
    assert actual == data
    result = recalculate(path, tmp_path / 'recalc', 'base', layout['outputs']['active_case'])
    sheet = openpyxl.load_workbook(result, data_only=True)['Decision']
    assert sheet['B9'].value == pytest.approx(1905.7439567)
    assert sheet['B17'].value == pytest.approx(677.8675480008182)
    assert sheet['B23'].value == pytest.approx(92.281)
    assert sheet['B24'].value == pytest.approx(1135.5954086991818)
    assert sheet['B38'].value == pytest.approx(103.631)
    assert sheet['C38'].value == pytest.approx(92.365)
    assert sheet['C40'].value == pytest.approx(92.365 / 1905.7439567)
    source = tmp_path / 'inputs.json'
    source.write_text(json.dumps(data))
    refreshed = tmp_path / 'refreshed.xlsx'
    refresh(path, source, refreshed)
    assert read_inputs(refreshed)[1]['presentation_version'] == 4
    data['other_assets'] = 92281000
    write_excel_model(data, tmp_path / 'included.xlsx', presentation_version=4)
    result = recalculate(tmp_path / 'included.xlsx', tmp_path / 'included', 'base', layout['outputs']['active_case'])
    assert openpyxl.load_workbook(result, data_only=True)['Decision']['B24'].value == 'N/A'
    modified = openpyxl.load_workbook(path)
    modified['Decision']['B24'] = '=0'
    modified.save(path)
    with pytest.raises(ValueError, match='changed'):
        read_inputs(path)


def test_decision_rejects_ambiguous_or_wrong_basis(tmp_path):
    data = inputs()
    data['historical'].append(copy.deepcopy(data['historical'][-1]))
    with pytest.raises(ValueError, match='unambiguous'):
        write_excel_model(data, tmp_path / 'duplicate.xlsx', presentation_version=4)
    data['historical'].pop()
    data['historical'][-1]['ownership_basis'] = 'whole_operation'
    with pytest.raises(ValueError, match='unambiguous'):
        write_excel_model(data, tmp_path / 'basis.xlsx', presentation_version=4)


def test_appended_inventory_selects_latest_and_never_falls_back_from_partial(tmp_path):
    data = inputs()
    prior = copy.deepcopy(data['historical'][-2:])
    for fact in prior:
        fact['period'] = 'FY2025'
    data['historical'].extend(prior)
    path = tmp_path / 'appended.xlsx'
    write_excel_model(data, path, presentation_version=4)
    sheet = openpyxl.load_workbook(path)['Decision']
    assert 'H1 2026' in sheet['A20'].value
    later = dict(prior[0], period='FY2026')
    data['historical'].append(later)
    with pytest.raises(ValueError, match='latest inventory'):
        write_excel_model(data, tmp_path / 'partial.xlsx', presentation_version=4)


def test_missing_quote_has_no_fabricated_market_comparison(tmp_path):
    if not shutil.which('soffice'):
        pytest.skip('LibreOffice needed')
    data = inputs()
    data['reference_price'] = None
    data['quote_date'] = None
    path = tmp_path / 'missing.xlsx'
    write_excel_model(data, path, presentation_version=4)
    _, layout, _, _ = read_inputs(path)
    result = recalculate(path, tmp_path / 'recalc', 'base', layout['outputs']['active_case'])
    sheet = openpyxl.load_workbook(result, data_only=True)['Decision']
    assert all(sheet[cell].value == 'N/A' for cell in ['B7', 'B9', 'B18', 'B24', 'B40', 'C41'])
