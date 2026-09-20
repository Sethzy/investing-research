"""Excel-first layout and numerical contracts, including native sensitivity formulas."""

import copy
import json
import shutil
import subprocess
from pathlib import Path

import openpyxl
import pytest

from investing_research.excel_layout import MARKER, write_excel_model
from investing_research.models import calculate_scenario


def read_manifest(book):
    return json.loads("".join(str(row[0].value) for row in list(book["_Model"].rows)[1:] if row[0].value))


def example(kind):
    return json.loads((Path(__file__).parents[1] / "examples" / f"model-{kind}.json").read_text())


def test_editable_cells_and_formula_integrity_manifest(tmp_path):
    data = example("mine")
    original = copy.deepcopy(data)
    path = tmp_path / "model.xlsx"
    write_excel_model(data, path)
    assert data == original
    book = openpyxl.load_workbook(path)
    manifest = read_manifest(book)
    assert book["_Model"]["A1"].value == MARKER
    assert book["_Model"].sheet_state == "hidden"
    assert manifest["baseline"] == data
    assert len(book["Summary"]._charts) == 2
    assert book["Assumptions"]["B4"].value == "base"
    for key in ("reference_price", "quote_date", "valuation_date"):
        assert "/" + key in manifest["inputs"]
    for qualified, formula in manifest["formula_cells"].items():
        sheet, cell = qualified.split("!")
        assert book[sheet][cell].value == formula
    assert len([key for key in manifest["outputs"] if key.startswith("sensitivity/")]) == 9
    assert book["Historical"].freeze_panes


def test_legacy_workbooks_remain_readable_after_presentation_upgrade(tmp_path):
    from investing_research.excel import read_inputs

    data = example("mine")
    old = tmp_path / "old.xlsx"
    new = tmp_path / "new.xlsx"
    write_excel_model(data, old, presentation_version=1)
    write_excel_model(data, new)
    old_inputs, old_layout, _, _ = read_inputs(old)
    new_inputs, new_layout, _, _ = read_inputs(new)
    assert old_inputs == new_inputs
    assert old_layout["formula_cells"] == new_layout["formula_cells"]
    assert old_layout["inputs"] == new_layout["inputs"]
    assert new_layout["presentation_version"] == 2


@pytest.mark.parametrize("kind,case", [("mine", "bear"), ("mine", "bull"), ("fcf", "base")])
def test_native_excel_matches_independent_engine(tmp_path, kind, case):
    soffice = shutil.which("soffice")
    if not soffice:
        pytest.skip("LibreOffice is needed for independent native formula evaluation")
    data = example(kind)
    # Exercise negative EBIT tax floor, closure, corporate expenses, ownership and
    # explicit attributable residual; unequal horizons catch phantom cash flows.
    data["scenarios"]["bear"]["years"] = data["scenarios"]["bear"]["years"][:1]
    if kind == "mine":
        for scenario in data["scenarios"].values():
            scenario["residual_value"] = 1234
            scenario["years"][0]["price"] = 1
            scenario["years"][-1]["closure"] = 4321
    path = tmp_path / "model.xlsx"
    write_excel_model(data, path)
    book = openpyxl.load_workbook(path)
    manifest = read_manifest(book)
    book["Assumptions"]["B4"] = case
    book.save(path)
    output = tmp_path / "recalculated"
    output.mkdir()
    subprocess.run(
        [
            soffice,
            f"-env:UserInstallation={tmp_path.joinpath('lo-profile').as_uri()}",
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(output),
            str(path),
        ],
        check=True,
        capture_output=True,
        timeout=90,
    )
    actual = openpyxl.load_workbook(output / path.name, data_only=True)

    def value(key):
        loc = manifest["outputs"][key]
        return actual[loc["sheet"]][loc["cell"]].value

    scenario = data["scenarios"][case]
    expected = calculate_scenario(data, scenario)
    for key in ("enterprise_value", "equity_value", "per_share", "minimum_illustrative_cash"):
        assert value(key) == pytest.approx(expected[key])
    for i, year in enumerate(expected["years"]):
        for key, wanted in year.items():
            assert (
                value(f"years/{i}/{key}") == pytest.approx(wanted)
                if wanted is not None
                else value(f"years/{i}/{key}") in (None, "")
            )
    for r, multiplier in enumerate((0.8, 1, 1.2)):
        for c, shift in enumerate((-0.02, 0, 0.02)):
            changed = copy.deepcopy(scenario)
            changed["discount_rate"] = max(0, scenario["discount_rate"] + shift)
            for year in changed["years"]:
                year["price" if kind == "mine" else "nopat"] *= multiplier
            assert value(f"sensitivity/{r}/{c}") == pytest.approx(
                calculate_scenario(data, changed)["per_share"]
            )
    for i in range(6):
        assert value(f"checks/{i}") == "PASS"
