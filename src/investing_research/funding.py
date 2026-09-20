"""Explicit financing overlay: funding gaps and dilution, not a new equity valuation."""

import math
import shutil
import tempfile
from pathlib import Path

import openpyxl
import xlsxwriter
from pydantic import Field

from .contracts import Record
from .excel import file_hash, recalculate, verify_snapshot
from .workspace import atomic_text, digest, now, read_json, write_json


VERSION = "1.0.0"


class FundingYear(Record):
    year: int
    interest_rate: float = Field(ge=0, le=1)
    debt_raised: float = Field(ge=0)
    debt_repaid: float = Field(ge=0)
    equity_raised: float = Field(ge=0)
    issue_price: float = Field(gt=0)
    dividends: float = Field(ge=0)
    minimum_cash: float = Field(ge=0)
    source: str = Field(min_length=1)
    rationale: str = Field(min_length=1)


class FundingPlan(Record):
    company_id: str
    model: str
    scenarios: dict[str, list[FundingYear]]
    rationale: str = Field(min_length=1)


def calculate(inputs, case, plan):
    if [y.year for y in plan] != [y["year"] for y in case["years"]]:
        raise ValueError("Funding periods must match the model scenario exactly")
    cash, debt, shares = (inputs[k] for k in ("cash", "debt", "diluted_shares"))
    initial_shares = shares
    rows = []
    for y, operating in zip(plan, case["years"]):
        if y.debt_repaid > debt + y.debt_raised:
            raise ValueError("Debt repayments exceed available debt")
        interest = debt * y.interest_rate
        # Residual proceeds are explicitly in model FCF; do not add them twice.
        closing_cash = (
            cash + operating["fcf"] - interest + y.debt_raised - y.debt_repaid + y.equity_raised - y.dividends
        )
        closing_debt = debt + y.debt_raised - y.debt_repaid
        closing_shares = shares + y.equity_raised / y.issue_price
        row = {
            "year": y.year,
            "opening_cash": cash,
            "model_fcf": operating["fcf"],
            "opening_debt": debt,
            "interest": interest,
            "closing_cash": closing_cash,
            "closing_debt": closing_debt,
            "closing_shares": closing_shares,
            "funding_gap": max(0, y.minimum_cash - closing_cash),
            "ownership_dilution": 1 - initial_shares / closing_shares,
        }
        rows.append(row)
        cash, debt, shares = closing_cash, closing_debt, closing_shares
    return rows


def build_workbook(path, inputs, model, plan, results):
    wb = xlsxwriter.Workbook(path, {"strings_to_formulas": False, "strings_to_urls": False})
    title = wb.add_format({"bold": True, "font_color": "white", "bg_color": "#142D43", "font_size": 16})
    label = wb.add_format({"text_wrap": True, "valign": "top"})
    numeric = wb.add_format({"num_format": '#,##0.00;(#,##0.00);"–"'})
    blue = wb.add_format({"num_format": '#,##0.00;(#,##0.00);"–"', "font_color": "#0000FF"})
    pct = wb.add_format({"num_format": "0.0%"})
    input_pct = wb.add_format({"num_format": "0.0%", "font_color": "#0000FF"})
    year_fmt = wb.add_format({"num_format": "0", "bold": True})
    red = wb.add_format({"font_color": "#A12232", "bg_color": "#FCE9EC"})
    readme = wb.add_worksheet("Readme")
    readme.hide_gridlines(2)
    readme.set_column("A:H", 18)
    readme.set_landscape()
    readme.fit_to_pages(1, 1)
    readme.print_area(0, 0, 31, 7)
    readme.merge_range("A1:H2", "Funding and dilution | " + plan.company_id, title)
    notes = [
        "This is a separate financing overlay linked to an immutable operating model. It is not a complete three-statement model or an adjusted price target.",
        "Currency: "
        + inputs["currency"]
        + ". Figures use absolute currency units. Debt/share inputs inherit the underlying model limitations; issued-share proxies are not verified dilution schedules.",
        "Interest uses opening debt. All new borrowing, repayment, equity issuance and dividends occur at period end. No interest tax shield, financing fees or intra-period liquidity is assumed.",
        "No automatic equity raise is invented to plug a shortfall. Positive funding gap means the supplied plan does not meet the cash floor. Negative cash is an unresolved financing need.",
        "Working capital, depreciation, capex and closure are inherited from the operating model. Do not subtract them again. No additional fair-value-per-share calculation is inferred from future financing.",
        "Edit the sourced funding-plan JSON and rerun funding-review to create a new snapshot. This workbook is a formula audit artifact; direct edits do not synchronize through excel-sync.",
    ]
    for i, note in enumerate(notes):
        readme.merge_range(5 + i * 4, 0, 7 + i * 4, 7, note, label)
    readme.write("A4", "Validation selector")
    readme.write("B4", "all")
    locations = {}
    keys = [
        "year",
        "opening_cash",
        "model_fcf",
        "opening_debt",
        "interest_rate",
        "interest",
        "debt_raised",
        "debt_repaid",
        "equity_raised",
        "issue_price",
        "dividends",
        "closing_cash",
        "closing_debt",
        "closing_shares",
        "minimum_cash",
        "funding_gap",
        "ownership_dilution",
    ]
    for case in ("bear", "base", "bull"):
        ws = wb.add_worksheet(case.title())
        ws.hide_gridlines(2)
        ws.set_column("A:A", 29)
        ws.set_column(1, len(results[case]), 20)
        ws.freeze_panes(4, 1)
        ws.merge_range(0, 0, 1, max(2, len(results[case])), case.title() + " | " + inputs["currency"], title)
        for row, key in enumerate(keys, 4):
            ws.write(row, 0, key.replace("_", " ").title(), label)
        for i, (y, out) in enumerate(zip(plan.scenarios[case], results[case]), 1):
            from xlsxwriter.utility import xl_col_to_name

            c, p = xl_col_to_name(i), xl_col_to_name(i - 1)
            formulas = {
                "opening_cash": f"={p}16" if i > 1 else None,
                "opening_debt": f"={p}17" if i > 1 else None,
                "interest": f"={c}8*{c}9",
                "closing_cash": f"={c}6+{c}7-{c}10+{c}11-{c}12+{c}13-{c}15",
                "closing_debt": f"={c}8+{c}11-{c}12",
                "closing_shares": f"={p}18+{c}13/{c}14" if i > 1 else f"=$B$24+{c}13/{c}14",
                "funding_gap": f"=MAX(0,{c}19-{c}16)",
                "ownership_dilution": f"=1-$B$24/{c}18",
            }
            values = {**y.model_dump(), **out}
            for r, key in enumerate(keys, 4):
                fmt = pct if key in ("interest_rate", "ownership_dilution") else numeric
                if formulas.get(key):
                    ws.write_formula(r, i, formulas[key], fmt)
                else:
                    ws.write(
                        r,
                        i,
                        values[key],
                        year_fmt if key == "year" else input_pct if key == "interest_rate" else blue,
                    )
                if key in out and key != "year":
                    locations[f"{case}/{i - 1}/{key}"] = (case.title(), r + 1, i + 1)
                ws.write_comment(
                    r, i, f"Source: {y.source}\nRationale: {y.rationale}\nOperating snapshot: {plan.model}"
                )
        ws.write("A24", "Initial diluted shares")
        ws.write("B24", inputs["diluted_shares"], blue)
        ws.conditional_format(
            19, 1, 19, len(results[case]), {"type": "cell", "criteria": ">", "value": 0, "format": red}
        )
        ws.set_landscape()
        ws.fit_to_pages(1, 1)
        ws.print_area(0, 0, 24, max(2, len(results[case])))
    wb.close()
    return locations


def review(ws, payload):
    plan = FundingPlan.model_validate(payload)
    ws.company(plan.company_id)
    path = ws.inside(plan.model)
    if path.name != "model.json":
        raise ValueError("Use a model.json snapshot")
    model = verify_snapshot(path)
    if model["company_id"] != plan.company_id:
        raise ValueError("Funding model belongs to another company")
    if set(plan.scenarios) != {"bear", "base", "bull"}:
        raise ValueError("Supply all three funding scenarios")
    inputs = read_json(path.parent / "inputs.json")
    results = {
        case: calculate(inputs, model["scenarios"][case], plan.scenarios[case]) for case in plan.scenarios
    }
    identity = digest({"version": VERSION, "plan": plan.model_dump(mode="json"), "model": model})
    folder = ws.root / "companies" / plan.company_id / "funding" / identity
    if folder.exists():
        old = read_json(folder / "funding.json")
        if (
            not old
            or old.get("id") != identity
            or old.get("scenarios") != results
            or old.get("plan") != plan.model_dump(mode="json")
            or not (folder / "model.xlsx").is_file()
            or file_hash(folder / "model.xlsx") != old["workbook_sha256"]
            or not (folder / "report.md").exists()
        ):
            raise ValueError("Funding snapshot damaged")
        return old
    folder.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=folder.parent, prefix=".funding-") as temp:
        staged = Path(temp) / "publish"
        staged.mkdir()
        workbook = staged / "model.xlsx"
        locations = build_workbook(workbook, inputs, model, plan, results)
        computed = recalculate(workbook, Path(temp) / "calc", "all", {"sheet": "Readme", "cell": "B4"})
        book = openpyxl.load_workbook(computed, data_only=True)
        try:
            for key, (sheet, row, col) in locations.items():
                case, index, field = key.split("/")
                actual = book[sheet].cell(row, col).value
                if type(actual) not in (int, float) or not math.isclose(
                    actual, results[case][int(index)][field], rel_tol=1e-9, abs_tol=1e-7
                ):
                    raise ValueError("Funding workbook parity failed: " + key)
        finally:
            book.close()
        shutil.copy2(computed, workbook)
        result = {
            "id": identity,
            "version": VERSION,
            "recorded_at": now(),
            "company_id": plan.company_id,
            "plan": plan.model_dump(mode="json"),
            "scenarios": results,
            "workbook_sha256": file_hash(workbook),
            "validation": "passed",
            "report": ws.relative(folder / "report.md"),
        }
        write_json(staged / "funding.json", result)
        from .reports import link

        lines = [
            f"# {plan.company_id} — funding and dilution",
            "",
            "**Synthetic example**"
            if model.get("synthetic")
            else "**Illustrative financing assumptions; underlying evidence limitations still apply.**",
            "",
            plan.rationale,
            "",
            f"Currency: {inputs['currency']}. [Operating snapshot]({link(path.parent / 'report.md', folder)})",
            "",
            "[Detailed financing workbook](model.xlsx) · [Plan and results](funding.json)",
            "",
            "Explicit end-period financing; opening-debt interest; no tax shield, fees or automatic balancing equity. Working capital/depreciation/capex are inherited from the operating snapshot. This is not a new equity valuation.",
            "",
        ]
        for case, rows in results.items():
            lines += [
                f"## {case.title()}",
                "",
                "| Year | Closing cash | Debt | Shares | Cash-floor shortfall | Ownership dilution |",
                "|---|---:|---:|---:|---:|---:|",
            ]
            lines += [
                f"| {r['year']} | {r['closing_cash']:,.2f} | {r['closing_debt']:,.2f} | {r['closing_shares']:,.2f} | {r['funding_gap']:,.2f} | {r['ownership_dilution']:.1%} |"
                for r in rows
            ]
            lines += [""]
        atomic_text(staged / "report.md", "\n".join(lines).rstrip() + "\n")
        staged.rename(folder)
    return result
