"""Excel input ownership, safe recalculation, immutable publication and source refresh."""

from __future__ import annotations

import copy
import hashlib
import json
import math
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree

import openpyxl

from .models import analyse, calculate_scenario, validate_model
from .workspace import Workspace, atomic_text, now, read_json, write_json


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def pointer_get(data, pointer):
    value = data
    for key in pointer.strip("/").split("/"):
        value = value[int(key)] if isinstance(value, list) else value.get(key)
    return value


def pointer_set(data, pointer, value):
    keys = pointer.strip("/").split("/")
    target = data
    for key in keys[:-1]:
        target = target[int(key)] if isinstance(target, list) else target[key]
    target[int(keys[-1]) if isinstance(target, list) else keys[-1]] = value


def _safe_package(path):
    if path.suffix.lower() != ".xlsx" or path.stat().st_size > 20_000_000:
        raise ValueError("Use a macro-free .xlsx under 20 MB")
    with zipfile.ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > 50_000_000:
            raise ValueError("Workbook expanded content exceeds 50 MB")
        for name in archive.namelist():
            if any(part in name.lower() for part in ("vbaproject", "externallinks/", "connections.xml", "embeddings/")):
                raise ValueError("External data, embedded objects and macros are not supported")
            if name.endswith(".rels"):
                for rel in ElementTree.fromstring(archive.read(name)):
                    if rel.attrib.get("TargetMode") == "External":
                        if not rel.attrib.get("Type", "").endswith("/hyperlink"):
                            raise ValueError("External workbook relationships are not supported")
                        if not rel.attrib.get("Target", "").startswith(("https://", "http://")):
                            raise ValueError("Only public-web citation hyperlinks are supported")


def manifest(book):
    if "_Model" not in book or book["_Model"]["A1"].value != "investing-excel-v1":
        raise ValueError("Use an Excel-first workbook created by invest excel-create")
    sheet = book["_Model"]
    if sheet.max_row > 500:
        raise ValueError("Workbook manifest is too large")
    chunks = [sheet.cell(row, 1).value for row in range(2, sheet.max_row + 1)]
    if not all(isinstance(chunk, str) for chunk in chunks):
        raise ValueError("Malformed workbook manifest")
    result = json.loads("".join(chunks))
    if result.get("schema_version") != 1:
        raise ValueError("Unsupported workbook schema")
    validate_model(result["baseline"])
    return result


def _load(path, *, values=False):
    return openpyxl.load_workbook(path, data_only=values, keep_links=False, keep_vba=False)


def _same_cell(left, right):
    # Excel/LibreOffice may remove optional quotes around simple sheet names.
    def normalized(value):
        if isinstance(value, str) and value.startswith("="):
            return re.sub(r"'([A-Za-z_][A-Za-z0-9_]*)'!", r"\1!", value)
        return value
    if type(left) in (int, float) and type(right) in (int, float):
        return math.isclose(left, right, rel_tol=1e-12, abs_tol=1e-12)
    return normalized(left) == normalized(right)


def read_inputs(path: Path):
    """Rebuild trusted cell maps from code, never trust a workbook's claimed formula list."""
    from .excel_layout import write_excel_model

    _safe_package(path)
    book = _load(path)
    try:
        declared = manifest(book)
        baseline = declared["baseline"]
        if max(len(s["years"]) for s in baseline["scenarios"].values()) > 50:
            raise ValueError("Excel models support at most 50 explicit annual periods")
        with tempfile.TemporaryDirectory(prefix="invest-excel-contract-") as temp:
            reference = Path(temp) / "reference.xlsx"
            write_excel_model(baseline, reference, presentation_version=declared.get("presentation_version", 1))
            expected_book = _load(reference)
            try:
                expected = manifest(expected_book)
                if declared != expected:
                    raise ValueError("Workbook metadata changed; regenerate the model after review")
                editable = {(v["sheet"], v["cell"]) for v in expected["inputs"].values()}
                selector = expected["outputs"]["active_case"]
                editable.add((selector["sheet"], selector["cell"]))
                if book.sheetnames != expected_book.sheetnames:
                    raise ValueError("Workbook sheets changed; model structure requires review")
                for sheet in book:
                    if sheet.max_row > 20_000 or sheet.max_column > 250:
                        raise ValueError("Workbook dimensions exceed the model boundary")
                    if sheet.title == "_Model":
                        continue
                    reference_sheet = expected_book[sheet.title]
                    for row in sheet:
                        for cell in row:
                            if (sheet.title, cell.coordinate) in editable:
                                if cell.data_type == "f":
                                    raise ValueError(f"Input {sheet.title}!{cell.coordinate} must be a value")
                            elif not _same_cell(cell.value, reference_sheet[cell.coordinate].value):
                                raise ValueError(f"Model formula/label changed at {sheet.title}!{cell.coordinate}; review and regenerate")
                    # Deleting trailing calculated rows must not evade comparison.
                    for row in reference_sheet:
                        for cell in row:
                            if cell.value is not None and (sheet.title, cell.coordinate) not in editable:
                                if not _same_cell(sheet[cell.coordinate].value, cell.value):
                                    raise ValueError(f"Model cell missing or changed at {sheet.title}!{cell.coordinate}")
            finally:
                expected_book.close()
        data = copy.deepcopy(baseline)
        changes = []
        for pointer, location in expected["inputs"].items():
            value = book[location["sheet"]][location["cell"]].value
            if pointer.endswith("date") and isinstance(value, (date, datetime)):
                value = (value.date() if isinstance(value, datetime) else value).isoformat()
            old = pointer_get(baseline, pointer)
            pointer_set(data, pointer, value)
            if value != old:
                changes.append({"path": pointer, "before": old, "after": value, **location})
                previous = data["provenance"].get(pointer, {})
                data["provenance"][pointer] = {
                    **previous,
                    "unit": previous.get("unit", "YYYY-MM-DD" if pointer.endswith("date") else data["currency"] + "/share"),
                    "source": f"workbook://{location['sheet']}/{location['cell']}",
                    "rationale": f"Explicit workbook override; prior value {old!r}. Prior source: {previous.get('source', 'unavailable')}",
                    "author": "workbook user",
                }
        case = book[selector["sheet"]][selector["cell"]].value
        if case not in ("bear", "base", "bull"):
            raise ValueError("Select bear, base or bull in the workbook")
        validate_model(data)
        return data, expected, case, changes
    finally:
        book.close()


def create(inputs: Path, output: Path, *, presentation_version: int = 2):
    from .excel_layout import write_excel_model

    if output.exists():
        raise ValueError("Output already exists; choose a new workbook filename")
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Output must be .xlsx")
    data = validate_model(read_json(inputs))
    output.parent.mkdir(parents=True, exist_ok=True)
    # Only publish complete packages; a failed writer must not occupy the target name.
    with tempfile.TemporaryDirectory(dir=output.parent, prefix=".excel-create-") as temp:
        staged = Path(temp) / "model.xlsx"
        write_excel_model(data, staged, presentation_version=presentation_version)
        read_inputs(staged)
        staged.rename(output)
    return {"workbook": str(output), "status": "editable_unverified", "next": "Run invest excel-sync on this workbook"}


def recalculate(source: Path, directory: Path, case, location):
    binary = shutil.which("soffice")
    if not binary:
        raise ValueError("LibreOffice/soffice is required to synchronize Excel; cached values are not accepted")
    incoming = directory / "incoming"
    outgoing = directory / "recalculated"
    incoming.mkdir(parents=True)
    outgoing.mkdir()
    path = incoming / "model.xlsx"
    book = _load(source)
    book[location["sheet"]][location["cell"]] = case
    book.save(path)
    book.close()
    try:
        process = subprocess.run(
            [binary, f"-env:UserInstallation={(directory / 'profile').as_uri()}", "--headless", "--convert-to", "xlsx", "--outdir", str(outgoing), str(path)],
            capture_output=True, timeout=90, check=False,
        )
    except (OSError, subprocess.TimeoutExpired):
        raise ValueError("Workbook recalculation failed or timed out; no result was published") from None
    result = outgoing / "model.xlsx"
    if process.returncode or not result.is_file():
        raise ValueError("Workbook recalculation failed; no result was published")
    return result


def _number(actual, expected, label):
    if isinstance(actual, bool) or not isinstance(actual, (int, float)) or not math.isfinite(actual):
        raise ValueError(f"Workbook result {label} is missing or invalid")
    if not math.isclose(actual, expected, rel_tol=1e-8, abs_tol=1e-7):
        raise ValueError(f"Workbook result {label} disagrees with independent calculation")
    return actual


def checked_results(path, layout, data, case):
    expected = calculate_scenario(data, data["scenarios"][case])
    result = copy.deepcopy(expected)
    book = _load(path, values=True)
    try:
        for sheet in book:
            if sheet.title != "_Model":
                for row in sheet:
                    if any(cell.data_type == "e" for cell in row):
                        raise ValueError(f"Excel calculation error on {sheet.title}")
        sensitivities = []
        for key, location in layout["outputs"].items():
            value = book[location["sheet"]][location["cell"]].value
            if key in ("enterprise_value", "equity_value", "per_share", "minimum_illustrative_cash"):
                result[key] = _number(value, expected[key], key)
            elif key.startswith("years/"):
                _, index, metric = key.split("/")
                if int(index) < len(expected["years"]):
                    number = expected["years"][int(index)].get(metric)
                    if number is not None:
                        result["years"][int(index)][metric] = _number(value, number, key)
            elif key.startswith("checks/") and value != "PASS":
                raise ValueError(f"Workbook check failed: {key}")
            elif key.startswith("sensitivity/"):
                _, r, c = key.split("/")
                multiplier_location = layout["outputs"][f"driver_multiplier/{r}"]
                rate_location = layout["outputs"][f"discount_rate/{c}"]
                multiplier = book[multiplier_location["sheet"]][multiplier_location["cell"]].value
                rate = book[rate_location["sheet"]][rate_location["cell"]].value
                expected_multiplier = (0.8, 1.0, 1.2)[int(r)]
                expected_rate = max(0, data["scenarios"][case]["discount_rate"] + (-0.02, 0, 0.02)[int(c)])
                _number(multiplier, expected_multiplier, "sensitivity driver")
                _number(rate, expected_rate, "sensitivity rate")
                changed = copy.deepcopy(data["scenarios"][case])
                changed["discount_rate"] = rate
                for year in changed["years"]:
                    year["price" if data["model_type"] == "finite_mine" else "nopat"] *= multiplier
                expected_value = calculate_scenario(data, changed)["per_share"]
                sensitivities.append({"driver_multiplier": multiplier, "discount_rate": rate, "per_share": _number(value, expected_value, key)})
        return result, sensitivities
    finally:
        book.close()


def synchronize(workspace: Workspace, workbook: Path):
    from .excel_reports import report

    workbook = workbook.resolve()
    if not workbook.is_relative_to(workspace.root):
        raise ValueError("Copy the working workbook into this workspace before synchronizing")
    data, layout, case, changes = read_inputs(workbook)
    company_id = data["company_id"]
    if not re.fullmatch(r"[a-z0-9]+-[a-z0-9.-]+", company_id):
        raise ValueError("Invalid company ID")
    source_hash = file_hash(workbook)
    result = analyse(data)  # independent checker; headline results below come from Excel.
    parent = workspace.root / "companies" / company_id / "models"
    parent.mkdir(parents=True, exist_ok=True)
    run_id = "excel-" + source_hash[:16]
    destination = parent / run_id
    if destination.exists():
        prior = read_json(destination / "model.json")
        if not prior or prior.get("source_workbook_sha256") != source_hash or not (destination / "report.md").is_file():
            raise ValueError("Existing workbook snapshot is incomplete; preserve it and choose a new workbook version")
        verify_snapshot(destination / "model.json")
        result = prior
    else:
        with tempfile.TemporaryDirectory(dir=parent, prefix=".excel-sync-") as temp:
            stage = Path(temp)
            published = stage / "published"
            published.mkdir()
            for selected in ("bear", "base", "bull"):
                recalculated = recalculate(workbook, stage / selected, selected, layout["outputs"]["active_case"])
                scenario, sensitivities = checked_results(recalculated, layout, data, selected)
                result["scenarios"][selected] = scenario
                if selected == "base":
                    result["sensitivity"] = sensitivities
                if selected == case:
                    shutil.copy2(recalculated, published / "model.xlsx")
            if file_hash(workbook) != source_hash:
                raise ValueError("Workbook changed during synchronization; retry after saving your edits")
            result.update(
                authority="excel", active_case=case, synchronized_at=now(),
                source_workbook_sha256=source_hash, workbook_sha256=file_hash(published / "model.xlsx"),
                source_workbook=workspace.relative(workbook), input_changes=changes,
                workbook_validation={"status": "passed", "engine": "LibreOffice", "checks": "Excel operating, valuation and sensitivity outputs matched independent Python calculations in all three cases"},
                run_id=run_id, output_dir=str(destination),
            )
            write_json(published / "inputs.json", data)
            write_json(published / "model.json", result)
            atomic_text(published / "report.md", report(data, result))
            published.rename(destination)
    write_json(workspace.root / "state/excel" / f"{company_id}.json", {
        "working_workbook": workspace.relative(workbook), "source_hash": source_hash,
        "model": workspace.relative(destination / "model.json"), "synchronized_at": result["synchronized_at"],
    })
    if workspace.settings_path.exists() and any(c.id == company_id for c in workspace.settings().companies):
        from .reports import dossier
        dossier(workspace, company_id)
    return result


def status(workspace: Workspace, company_id=None):
    entries = []
    for path in sorted((workspace.root / "state/excel").glob("*.json")):
        if company_id and path.stem != company_id:
            continue
        record = read_json(path)
        working = workspace.inside(record["working_workbook"])
        current = file_hash(working) if working.is_file() else None
        entries.append({"company_id": path.stem, **record, "status": "current" if current == record["source_hash"] else "unsynchronized_edits" if current else "working_workbook_missing"})
    return entries


def verify_snapshot(model_path: Path):
    """Validate immutable workbook/result binding before allowing an actionable view."""
    result = read_json(model_path)
    data = read_json(model_path.parent / "inputs.json")
    workbook = model_path.parent / "model.xlsx"
    if not result or not data or not workbook.is_file():
        raise ValueError("Excel snapshot is incomplete")
    if file_hash(workbook) != result.get("workbook_sha256"):
        raise ValueError("Snapshot workbook was edited; synchronize a working copy first")
    expected = analyse(data)

    def compare(actual, reference):
        if isinstance(reference, dict):
            return isinstance(actual, dict) and all(k in actual and compare(actual[k], v) for k, v in reference.items())
        if isinstance(reference, list):
            return isinstance(actual, list) and len(actual) == len(reference) and all(compare(a, b) for a, b in zip(actual, reference))
        if type(reference) in (int, float):
            return type(actual) in (int, float) and math.isclose(actual, reference, rel_tol=1e-8, abs_tol=1e-7)
        return actual == reference

    if not compare(result, expected) or result.get("workbook_validation", {}).get("status") != "passed":
        raise ValueError("Excel results do not reproduce from their accepted assumptions")
    parsed, _, _, _ = read_inputs(workbook)
    # Provenance gains a deterministic override record at import; financial inputs must match.
    if not compare({k: v for k, v in parsed.items() if k != "provenance"}, {k: v for k, v in data.items() if k != "provenance"}):
        raise ValueError("Excel snapshot assumptions differ from saved inputs")
    return result


def validate_snapshot_live(model_path: Path):
    verify_snapshot(model_path)
    workbook = model_path.parent / "model.xlsx"
    data, layout, _, _ = read_inputs(workbook)
    with tempfile.TemporaryDirectory(prefix="invest-excel-revalidate-") as directory:
        for case in ("bear", "base", "bull"):
            recalculated = recalculate(workbook, Path(directory) / case, case, layout["outputs"]["active_case"])
            checked_results(recalculated, layout, data, case)
    return {"status": "passed", "engine": "LibreOffice", "checks": "All scenarios, operating outputs and sensitivities independently recalculated"}


def refresh(workbook: Path, inputs: Path, output: Path):
    from .excel_layout import write_excel_model

    if output.exists() or output.suffix.lower() != ".xlsx":
        raise ValueError("Choose a new .xlsx output; refresh never overwrites the working model")
    edited, layout, case, _ = read_inputs(workbook)
    old = layout["baseline"]
    incoming = validate_model(read_json(inputs))
    if any(old.get(key) != incoming.get(key) for key in ("company_id", "currency", "model_type", "purpose", "synthetic")):
        raise ValueError("Source refresh must retain company, currency, model type and evidence purpose")
    if any([y["year"] for y in old["scenarios"][s]["years"]] != [y["year"] for y in incoming["scenarios"][s]["years"]] for s in old["scenarios"]):
        raise ValueError("Forecast horizon changes require an explicit new model and assumption review")
    for name in old["scenarios"]:
        for prior, proposed in zip(old["scenarios"][name]["years"], incoming["scenarios"][name]["years"]):
            if prior.get("price_currency") != proposed.get("price_currency"):
                raise ValueError("Price-currency changes require a reviewed model conversion")
    merged = copy.deepcopy(incoming)
    conflicts = []
    preserved = []
    quote_overridden = any(pointer_get(edited, p) != pointer_get(old, p) or old["provenance"].get(p, {}).get("author") == "workbook user" for p in ("/reference_price", "/quote_date"))
    for pointer in layout["inputs"]:
        previous, user, source = (pointer_get(d, pointer) for d in (old, edited, incoming))
        existing_override = old["provenance"].get(pointer, {}).get("author") == "workbook user"
        if user != previous or existing_override or (quote_overridden and pointer in ("/reference_price", "/quote_date")):
            pointer_set(merged, pointer, user)
            provenance = copy.deepcopy(edited["provenance"].get(pointer, {}))
            merged["provenance"][pointer] = {
                **provenance,
                "unit": provenance.get("unit", "YYYY-MM-DD" if pointer.endswith("date") else merged["currency"] + "/share"),
                "source": provenance.get("source", "workbook://paired-quote"),
                "rationale": provenance.get("rationale", "Quote date retained with the user's price override"),
                "author": "workbook user", "refresh_source_value": source,
            }
            preserved.append(pointer)
            previous_source = old["provenance"].get(pointer, {}).get("refresh_source_value", previous)
            if source != previous_source and source != user:
                conflicts.append({"path": pointer, "prior_source": previous_source, "new_source": source, "kept_user_value": user})
    # Historical source versions remain available rather than silently replacing old facts.
    historical = {json.dumps(row, sort_keys=True): row for row in old["historical"] + incoming["historical"]}
    merged["historical"] = list(historical.values())
    validate_model(merged)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=output.parent, prefix=".excel-refresh-") as temp:
        staged = Path(temp) / "model.xlsx"
        write_excel_model(merged, staged, presentation_version=max(2, layout.get("presentation_version", 1)))
        book = _load(staged)
        location = manifest(book)["outputs"]["active_case"]
        book[location["sheet"]][location["cell"]] = case
        book.save(staged)
        book.close()
        read_inputs(staged)
        staged.rename(output)
    receipt = {"status": "refresh_needs_review" if conflicts else "refreshed_unverified", "workbook": str(output), "preserved_overrides": preserved, "conflicts": conflicts}
    write_json(output.with_suffix(".refresh.json"), receipt)
    return receipt
