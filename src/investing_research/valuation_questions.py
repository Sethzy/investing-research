"""Conditional cash/value hurdles; no automatic rating or financing assumptions.

Use with: python -m investing_research.valuation_questions INPUT OUTPUT_DIRECTORY
All money and shares are in millions, in one declared currency.
"""
from __future__ import annotations

import argparse
import json
import math
from pathlib import Path

import xlsxwriter


def annuity(rate, years, first_year=1):
    if not math.isfinite(rate) or rate < 0:
        raise ValueError('rate must be finite and nonnegative')
    if any(isinstance(v, bool) or not isinstance(v, int) or v < 1 for v in (years, first_year)):
        raise ValueError('years and first_year must be positive integers')
    return sum((1 + rate) ** -t for t in range(first_year, first_year + years))


def project_hurdle(residual, capital, ownership, rate, years, first_year, closure):
    if not all(math.isfinite(x) for x in (residual, capital, ownership, closure)):
        raise ValueError('nonfinite project input')
    if capital < 0 or closure < 0 or not 0 < ownership <= 1:
        raise ValueError('invalid capital, closure or ownership')
    factor = annuity(rate, years, first_year)
    liability = closure / (1 + rate) ** (first_year + years - 1)
    return (max(residual, 0) / ownership + capital + liability) / factor


ANNUAL = ('ocf', 'working_capital_movement', 'interest_income', 'extension_fee',
          'interest_expense', 'depreciation', 'equipment', 'mine_development', 'exploration')
DRIVERS = ('price', 'shares', 'cash', 'debt', 'investments', 'investment_haircut', 'closure',
           'life', 'rate', 'annual_cash', 'ownership', 'project_capital', 'project_life',
           'first_cash_year', 'project_closure', 'cash_buffer', 'other_commitments', 'tax_rate')


def validate(data):
    if set(data) != {'company_id', 'company', 'currency', 'as_of', 'annual', 'drivers', 'provenance', 'limitations'}:
        raise ValueError('Use the documented input contract')
    if not data['company_id'] or not data['company'] or not data['currency'] or not data['limitations']:
        raise ValueError('Company, currency and limitations required')
    from datetime import date
    date.fromisoformat(data['as_of'])
    if set(data['drivers']) != set(DRIVERS) or len(data['annual']) != 2:
        raise ValueError('Two annual periods and all drivers required')
    records = [(f'drivers.{k}', v) for k, v in data['drivers'].items()]
    periods = []
    for i, record in enumerate(data['annual']):
        if set(record) != set(ANNUAL) | {'period'}:
            raise ValueError('Incomplete annual period')
        periods.append(record['period'])
        records += [(f'annual.{i}.{k}', record[k]) for k in ANNUAL]
    if periods != sorted(set(periods)):
        raise ValueError('Use distinct chronological annual periods')
    for key, value in records:
        if isinstance(value, bool) or not isinstance(value, (float, int)) or not math.isfinite(value):
            raise ValueError(f'Nonfinite input: {key}')
        if value < 0 and not key.endswith(('ocf', 'working_capital_movement')):
            raise ValueError(f'Negative input: {key}')
        source = data['provenance'].get(key, {})
        if source.get('status') not in {'reported', 'assumption', 'derived'} or not all(
            isinstance(source.get(k), str) and source[k].strip() for k in ('source', 'basis', 'rationale')
        ):
            raise ValueError(f'Missing provenance: {key}')
    d = data['drivers']
    if d['shares'] <= 0 or d['price'] <= 0 or not 0 < d['ownership'] <= 1:
        raise ValueError('Positive quote, shares and ownership required')
    if not 0 <= d['investment_haircut'] <= 1 or not 0 <= d['tax_rate'] <= 1:
        raise ValueError('Haircut and tax rate must be fractions')
    annuity(d['rate'], d['life'])
    annuity(d['rate'], d['project_life'], d['first_cash_year'])


def calculate(data):
    validate(data)
    d = data['drivers']
    cash = []
    for a in data['annual']:
        before = a['ocf'] - a['working_capital_movement'] - a['interest_income'] - a['extension_fee'] + a['interest_expense']
        tax = max(0, before - a['depreciation']) * d['tax_rate']
        cash.append(before - tax - a['equipment'] - a['mine_development'] - a['exploration'])
    net_assets = d['cash'] - d['debt'] + d['investments'] * (1 - d['investment_haircut'])
    closing = d['closure'] / (1 + d['rate']) ** d['life']
    factor = annuity(d['rate'], d['life'])
    equity = net_assets + d['annual_cash'] * factor - closing
    market = d['price'] * d['shares']
    residual = market - equity
    available = max(0, d['cash'] - d['cash_buffer'] - d['other_commitments'] - d['debt'])
    kwargs = dict(capital=d['project_capital'], ownership=d['ownership'], rate=d['rate'],
                  years=d['project_life'], first_year=d['first_cash_year'], closure=d['project_closure'])
    return dict(adjusted_cash=cash, equity=equity, per_share=equity / d['shares'], residual=residual,
                required_operating_cash=(market - net_assets + closing) / factor,
                project_break_even=project_hurdle(0, **kwargs),
                project_market_hurdle=project_hurdle(residual, **kwargs),
                project_delay_hurdle=project_hurdle(residual, **dict(kwargs, first_year=d['first_cash_year'] + 1)),
                cash_available=available, maximum_whole_capital=available / d['ownership'],
                funding_gap=max(0, d['project_capital'] * d['ownership'] - available),
                existing_cash_shortfall=max(0, d['cash_buffer'] + d['other_commitments'] + d['debt'] - d['cash']))


def write_workbook(data, path):
    """Maintained companion template; intentionally separate from accepted operating forecasts."""
    validate(data)
    if Path(path).exists():
        raise ValueError('Preserve earlier workbook; choose a new path')
    book = xlsxwriter.Workbook(path)
    common = {'font_name': 'Arial', 'font_size': 11, 'valign': 'vcenter'}
    fmt = {k: book.add_format(common | v) for k, v in {
        'title': {'bold': True, 'font_size': 17, 'bg_color': '#17364D', 'font_color': '#FFFFFF'},
        'label': {'text_wrap': True}, 'note': {'text_wrap': True, 'valign': 'top', 'font_color': '#536575', 'font_size': 10},
        'price': {'num_format': '0.00', 'font_color': '#0000FF'},
        'price_total': {'num_format': '0.00', 'bold': True, 'bg_color': '#E7F1F3', 'top': 1},
        'num': {'num_format': '#,##0.0;(#,##0.0);"-"'},
        'input': {'num_format': '#,##0.0;(#,##0.0);"-"', 'font_color': '#0000FF'},
        'link': {'num_format': '#,##0.0;(#,##0.0);"-"', 'font_color': '#008000'},
        'pct': {'num_format': '0.0%', 'font_color': '#0000FF'},
        'pct_calc': {'num_format': '0.0%'},
        'pct_link': {'num_format': '0.0%', 'font_color': '#008000'},
        'total': {'num_format': '#,##0.0;(#,##0.0);"-"', 'bold': True, 'bg_color': '#E7F1F3', 'top': 1},
    }.items()}
    for name, title in [('Value hurdle', 'What must the operating business earn?'),
                        ('Cash quality', 'Cash generation, with timing removed'),
                        ('Project hurdle', 'What would the omitted project need to deliver?')]:
        s = book.add_worksheet(name)
        s.hide_gridlines(2)
        s.set_zoom(90)
        s.set_column('A:A', 53)
        s.set_column('B:C', 19)
        s.set_column('D:D', 21)
        s.set_default_row(22)
        s.freeze_panes(5, 1)
        s.merge_range('A1:D2', title, fmt['title'])
        s.merge_range('A3:D4', f"{data['company']} | {data['as_of']} | {data['currency']} millions, except price, years and rates. Conditional tests; not a price target.", fmt['note'])
        s.set_portrait()
        s.set_paper(9)
        s.set_margins(.3, .3, .35, .35)
        s.set_print_scale(80)
        s.set_footer('&L' + name + '&RPage &P of &N', {'margin': .18})
    d = data['drivers']
    v = book.get_worksheet_by_name('Value hurdle')
    c = book.get_worksheet_by_name('Cash quality')
    p = book.get_worksheet_by_name('Project hurdle')

    def line(s, row, label, value, kind='num', col=1, key=None):
        s.write(row - 1, 0, label, fmt['label'])
        if isinstance(value, str) and value.startswith('='):
            if kind == 'pct':
                kind = 'pct_link' if '!' in value else 'pct_calc'
            elif kind == 'num' and '!' in value:
                kind = 'link'
            s.write_formula(row - 1, col, value, fmt[kind])
        else:
            s.write(row - 1, col, value, fmt[kind])
        if key:
            info = data['provenance'][key]
            s.write_comment(row - 1, col, f"{info['status']}: {info['basis']}\n{info['source']}\n{info['rationale']}")

    def driver(s, row, label, key, kind='input'):
        line(s, row, label, d[key], kind, key='drivers.' + key)

    def note(s, first, last, text):
        s.merge_range(f'A{first}:D{last}', text, fmt['note'])

    for col, a in enumerate(data['annual'], 1):
        x = 'BC'[col - 1]
        c.write(4, col, a['period'], fmt['label'])
        for row, field, label in [(6,'ocf','Reported operating cash'),(7,'working_capital_movement','Working-capital / provision movements'),
            (9,'interest_income','Remove accrued investment interest'),(10,'extension_fee','Remove note extension fee'),
            (11,'interest_expense','Add financing interest back'),(13,'depreciation','Depreciation and amortisation'),
            (17,'equipment','Equipment cash spending'),(18,'mine_development','Mine-development cash spending'),
            (19,'exploration','Exploration cash spending')]:
            line(c,row,label,a[field],'input',col,f'annual.{col-1}.{field}')
        for row,label,formula in [(8,'Cash before movements',f'={x}6-{x}7'),
            (12,'Operating cash before tax / investment',f'={x}8-{x}9-{x}10+{x}11'),
            (14,'Taxable-profit approximation',f'={x}12-{x}13'),
            (16,'Tax at assumed rate',f'=MAX(0,{x}14)*{x}15'),
            (20,'Historically adjusted operating cash',f'={x}12-{x}16-SUM({x}17:{x}19)')]:
            line(c,row,label,formula,'total' if row in (12,20) else 'num',col)
        if col == 1:
            driver(c,15,'Assumed tax rate','tax_rate','pct')
        else:
            line(c,15,'Assumed tax rate','=B15','pct',col)
    line(c,22,'Two-period average','=AVERAGE(B20:C20)','total')
    note(c,24,27,'These are adjusted historical cash flows, not through-cycle or forward forecasts. Working capital is held flat; tax uses accounting depreciation as an approximation. All equipment, development and exploration spending is deducted, without inventing a sustaining/growth split. Corporate costs remain included.')
    note(c,29,32,'Investment interest and the note fee are removed so investment principal can be assessed separately. Financing interest is added back: this is cash before financing, with debt deducted separately in the value bridge. Lease principal is financing. No new lease additions are assumed.')
    note(c,34,37,'Sources and reasoning are in input cell notes. Annual periods are not divided into artificial quarters. The recent strong half-year is not annualized. Taxes, royalty exposure, commodity cycle and future capital needs still require forward estimates.')
    c.print_area('A1:D38')
    for row,label,key in [(6,'Reference price per share','price'),(7,'Issued share denominator, millions','shares'),
        (10,'Corporate cash','cash'),(11,'Debt / hire purchase','debt'),(12,'Investment carrying amounts','investments'),
        (13,'Assumed haircut to carrying amounts','investment_haircut'),(16,'Assumed closure cash, attributable','closure'),
        (17,'Assumed remaining operating years','life'),(18,'Assumed discount rate','rate'),(19,'Annual cash, attributable after corporate costs','annual_cash')]:
        driver(v,row,label,key,'pct' if key in ('investment_haircut','rate') else 'price' if key == 'price' else 'input')
    for row,label,formula in [(8,'Market equity value','=B6*B7'),(14,'Conditional investment value','=B12*(1-B13)'),
        (20,'Finite cash-flow present-value factor','=IF(B18=0,B17,(1-(1+B18)^(-B17))/B18)'),
        (21,'Operating value after closure','=B19*B20-B16/(1+B18)^B17'),(23,'Net cash and conditional investments','=B10-B11+B14'),
        (24,'Conditional equity value','=B21+B23'),(25,'Conditional value per share','=B24/B7'),
        (26,'Market price value still unexplained','=B8-B24'),(28,'Annual cash required by reference price','=(B8-B23+B16/(1+B18)^B17)/B20')]:
        line(v,row,label,formula,'price_total' if row == 25 else 'total' if row in (24,26,28) else 'num')
    for row in (19, 28):
        v.set_row(row-1, 32)
    note(v,30,32,'Annual cash is an explicit stress input, not a forecast or the historical average. First payment is one year after valuation; no stub period, ramp or terminal value. Closure is a nominal assumed payment at the end of the chosen life. No accounting closure provision is added again.')
    v.set_h_pagebreaks([33])
    v.merge_range('A34:D35','Conditional value per share across cash and life',fmt['title'])
    v.write('A37','Years / annual cash',fmt['label'])
    for col,cash in enumerate([50,100,150],1):
        v.write(36,col,cash,fmt['input'])
    for row,life in enumerate([5,8,10],38):
        v.write(row-1,0,life,fmt['input'])
        for col in range(1,4):
            x='BCD'[col-1]
            v.write_formula(row-1,col,f'=($B$23+{x}$37*IF($B$18=0,$A{row},(1-(1+$B$18)^(-$A{row}))/$B$18)-$B$16/(1+$B$18)^$A{row})/$B$7',fmt['price_total'])
    note(v,42,45,'Matrix axes are analyst stress tests, not management guidance or probability-weighted cases. Change cash and life together when geology, disruption or investment requires it. Do not describe the highest value as a price target. Investment haircut also remains an assumption.')
    note(v,47,50,'The residual can reflect omitted assets, longer or stronger cash generation, stale balance-sheet figures, different required returns, or market optimism. It cannot be attributed exclusively to the development project. No personal margin-of-safety threshold is assumed.')
    v.print_area('A1:D51')
    line(p,6,'Residual from conditional operating case',"='Value hurdle'!B26",'link')
    for row,label,key in [(7,'Project ownership','ownership'),(8,'Whole-project capital, paid today','project_capital'),
        (9,'Whole-project operating years','project_life'),(10,'First cash receipt: years from today','first_cash_year'),
        (12,'Whole-project closure cash','project_closure'),(20,'Additional corporate cash buffer','cash_buffer'),
        (21,'Other reserved commitments','other_commitments')]:
        driver(p,row,label,key,'pct' if key=='ownership' else 'input')
    line(p,11,'Discount rate',"='Value hurdle'!B18",'pct')
    for row,label,formula in [(13,'Delayed cash-flow PV factor','=IF(B11=0,B9,(1-(1+B11)^(-B9))/B11)/(1+B11)^(B10-1)'),
        (14,'Present value of project closure','=B12/(1+B11)^(B10+B9-1)'),(15,'Attributable capital payment','=B8*B7'),
        (16,'Whole-project annual cash: economic break-even','=(B8+B14)/B13'),
        (17,'Whole-project annual cash: cover residual too','=(MAX(0,B6)/B7+B8+B14)/B13'),
        (19,'Corporate cash available before reserves',"='Value hurdle'!B10"),
        (22,'Debt reserved',"='Value hurdle'!B11"),
        (23,'Cash available for project','=MAX(0,B19-B20-B21-B22)'),
        (24,'Maximum whole-project capital from this cash','=B23/B7'),
        (25,'Unfunded attributable capital','=MAX(0,B15-B23)'),
        (26,'Existing cash shortfall before project','=MAX(0,B20+B21+B22-B19)'),
        (28,'Residual hurdle: one-year later receipts','=(MAX(0,B6)/B7+B8+B14/(1+B11))/(B13/(1+B11))'),
        (29,'Funding gap if capital is 25% higher','=MAX(0,B15*1.25-B23)')]:
        line(p,row,label,formula,'total' if row in (16,17,23,24,25,28,29) else 'num')
    for row in (16,17,24,28,29):
        p.set_row(row-1,34)
    note(p,31,34,'Cash is after tax, sustaining capital and working capital, before financing, at whole-project level. No production or metal-price target is inferred. Capital is assumed paid immediately; later receipts include delay but no ramp. The market-residual hurdle is conditional, not the project value.')
    note(p,36,39,'Funding excludes future operating inflows, dividends, asset sales and new debt/equity. The buffer is additional to reserved commitments. Check overlaps before reuse. A cash gap is not an approved financing plan. Economic break-even and liquidity are different questions.')
    p.print_area('A1:D40')
    book.close()


def verify_companion(directory):
    """Verify source/formula contract and independently recalculated cached results."""
    import tempfile
    import openpyxl
    directory = Path(directory)
    data = json.loads((directory / 'inputs.json').read_text())
    expected_results = calculate(data)
    if json.loads((directory / 'results.json').read_text()) != expected_results:
        raise ValueError('Companion calculated results changed')
    actual = openpyxl.load_workbook(directory / 'questions.xlsx')
    with tempfile.TemporaryDirectory() as temp:
        expected_path = Path(temp) / 'expected.xlsx'
        write_workbook(data, expected_path)
        expected = openpyxl.load_workbook(expected_path)
        if actual.sheetnames != expected.sheetnames:
            raise ValueError('Companion sheets changed')
        for sheet in expected:
            target = actual[sheet.title]
            for row in target:
                for cell in row:
                    if cell.value != sheet[cell.coordinate].value:
                        raise ValueError(f'Companion input/formula changed: {sheet.title}!{cell.coordinate}')
            for row in sheet:
                for cell in row:
                    if cell.value != target[cell.coordinate].value:
                        raise ValueError(f'Companion input/formula missing: {sheet.title}!{cell.coordinate}')
    cached = openpyxl.load_workbook(directory / 'questions.xlsx', data_only=True)
    checks = {'Cash quality': {'B20': expected_results['adjusted_cash'][0], 'C20': expected_results['adjusted_cash'][1]},
              'Value hurdle': {'B24': expected_results['equity'], 'B25': expected_results['per_share'],
                               'B26': expected_results['residual'], 'B28': expected_results['required_operating_cash']},
              'Project hurdle': {'B16': expected_results['project_break_even'], 'B17': expected_results['project_market_hurdle'],
                                 'B23': expected_results['cash_available'], 'B24': expected_results['maximum_whole_capital'],
                                 'B25': expected_results['funding_gap'], 'B26': expected_results['existing_cash_shortfall'],
                                 'B28': expected_results['project_delay_hurdle']}}
    d = data['drivers']
    net = d['cash'] - d['debt'] + d['investments'] * (1 - d['investment_haircut'])
    for row, life in enumerate([5, 8, 10], 38):
        for col, cash in zip('BCD', [50, 100, 150]):
            checks['Value hurdle'][f'{col}{row}'] = (net + cash * annuity(d['rate'], life) - d['closure'] / (1+d['rate'])**life) / d['shares']
    for sheet in cached:
        for row in sheet:
            for cell in row:
                if cell.data_type == 'e' or (actual[sheet.title][cell.coordinate].data_type == 'f' and cell.value is None):
                    raise ValueError(f'Uncalculated or error cell {sheet.title}!{cell.coordinate}')
    for name, cells in checks.items():
        for cell, expected in cells.items():
            value = cached[name][cell].value
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isclose(value, expected, rel_tol=1e-9, abs_tol=1e-7):
                raise ValueError(f'Companion numerical mismatch: {name}!{cell}')
    return {'status': 'pass', 'sheets': actual.sheetnames, 'checked_outputs': sum(map(len, checks.values()))}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    data = json.loads(args.input.read_text())
    results = calculate(data)
    args.output.mkdir(parents=True, exist_ok=False)
    (args.output / 'inputs.json').write_text(json.dumps(data, indent=2) + '\n')
    (args.output / 'results.json').write_text(json.dumps(results, indent=2) + '\n')
    write_workbook(data, args.output / 'questions.xlsx')
    import shutil
    import tempfile
    from .excel import recalculate
    with tempfile.TemporaryDirectory() as temp:
        calculated = recalculate(args.output / 'questions.xlsx', Path(temp), data['drivers']['price'],
                                 {'sheet': 'Value hurdle', 'cell': 'B6'})
        shutil.copyfile(calculated, args.output / 'questions.xlsx')
    verify_companion(args.output)


if __name__ == '__main__':
    main()
